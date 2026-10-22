#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
================================================================================
SISTEMA DE MONITOREO AUTOMATIZADO DE HORAS EN PLANTAS ELÉCTRICAS (GRUPOS ELECTRÓGENOS)
AIRTEK ENERGÍA NODOS - 100% GRATUITO
================================================================================
- Descarga desatendida desde enlace público de OneDrive Online (sin API de pago).
- Lectura de hoja 'Lista y datos de G.E.' con normalización universal de duraciones.
- Evaluación de umbral operativo (por defecto >= 40 horas).
- Generación y despacho de alertas:
    a) Correo Electrónico HTML (SMTP gratuito: Gmail / Outlook / Office365)
    b) Mensajes directos en Telegram Bot (100% gratuito)
- Memoria de estado anti-spam: evita saturar con alertas repetidas en cada ciclo.
"""

import os
import sys
import re
import json
import logging
import smtplib
from datetime import datetime, timedelta, time
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

import requests
import openpyxl

# Forzar codificación UTF-8 en consola Windows
if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

# ==============================================================================
# CONFIGURACIÓN GENERAL Y ENLACE DE ONEDRIVE
# ==============================================================================
ONEDRIVE_SHARING_URL = os.getenv(
    "ONEDRIVE_URL",
    "https://1drv.ms/x/c/2d21d74df8d2b570/IQBSDwn1pXdxTr7dW-quR5Y8AWR8dWvnipvRTS9CM_LKU-4?e=6NFapu"
)

LOCAL_EXCEL_FILENAME = "registro_monitoreo_plantas_live.xlsx"
STATE_FILE = "plantas_alerta_state.json"
LOG_FILE = "monitor_plantas.log"

UMBRAL_HORAS_DEFAULT = 40.0
INCREMENTO_MIN_ALERTA = 6.0   # Solo re-alertar si las horas aumentaron >= 6h desde el último aviso
EXPIRACION_ESTADO_HORAS = 24  # O si han pasado más de 24h desde el último aviso

# ==============================================================================
# CREDENCIALES DE NOTIFICACIÓN (Variables de entorno o valores por defecto)
# ==============================================================================
# 1. Telegram Bot (100% gratuito)
# Obtén tu token gratis hablando con @BotFather en Telegram
TELEGRAM_BOT_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN", "")

# Mapeo de destinatarios de Telegram por región o correo (chat_id personal o de grupo privado)
# Para obtener tu chat_id, inicia chat con @userinfobot en Telegram
TELEGRAM_DEFAULT_CHAT_ID = os.getenv("TELEGRAM_CHAT_ID", "")

# 2. Correo SMTP (100% gratuito con Gmail o Outlook)
SMTP_SERVER = os.getenv("SMTP_SERVER", "smtp.gmail.com")
SMTP_PORT = int(os.getenv("SMTP_PORT", "587"))
SMTP_USER = os.getenv("SMTP_USER", "")
SMTP_PASS = os.getenv("SMTP_PASS", "")  # En Gmail usar "Contraseña de aplicación" de 16 letras
SMTP_FROM = os.getenv("SMTP_FROM", SMTP_USER)

# Configuración de Logging
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.FileHandler(LOG_FILE, encoding='utf-8'),
        logging.StreamHandler(sys.stdout)
    ]
)
logger = logging.getLogger("MonitorPlantas")


# ==============================================================================
# MÓDULO 1: DESCARGA GRATUITA Y DIRECTA DESDE ONEDRIVE ONLINE
# ==============================================================================
def download_onedrive_excel(sharing_url: str, output_path: str) -> bool:
    """
    Descarga el archivo Excel en tiempo real desde un link compartido de OneDrive
    sin credenciales pagas, navegando la sesión pública de Microsoft.
    """
    logger.info("Conectando con OneDrive Online para obtener última versión...")
    session = requests.Session()
    session.headers.update({
        'User-Agent': (
            'Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
            'AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36'
        )
    })

    try:
        res = session.get(sharing_url, allow_redirects=True, timeout=25)
        if res.status_code != 200:
            logger.error(f"Error accediendo a OneDrive: HTTP {res.status_code}")
            return False

        # Extraer URL de descarga directa generada dinámicamente por OneDrive
        dl_matches = re.findall(r'\"(https://[^\"]*download\.aspx\?[^\"]*)\"', res.text)
        if not dl_matches:
            logger.error("No se localizó el enlace de descarga en la página de OneDrive.")
            return False

        download_url = None
        for u in dl_matches:
            u_clean = u.replace(r'\u0026', '&')
            if 'onedrive.live.com' in u_clean:
                download_url = u_clean
                break
            elif not download_url:
                download_url = u_clean

        dl_res = session.get(download_url, allow_redirects=True, timeout=45)
        if dl_res.status_code == 200 and dl_res.content[:4] == b'PK\x03\x04':
            with open(output_path, 'wb') as f:
                f.write(dl_res.content)
            size_mb = len(dl_res.content) / (1024 * 1024)
            logger.info(f"Excel descargado exitosamente: {output_path} ({size_mb:.2f} MB)")
            return True
        else:
            logger.error(f"Fallo al descargar binario: HTTP {dl_res.status_code}")
            return False

    except Exception as e:
        logger.error(f"Excepción descargando Excel de OneDrive: {e}")
        return False


# ==============================================================================
# MÓDULO 2: PARSEADOR UNIVERSAL DE DURACIONES EN EXCEL
# ==============================================================================
def parse_excel_duration_to_hours(val) -> float:
    """
    Convierte cualquier representación interna de Excel a horas decimales exactas:
    - timedelta (ej: 3 days, 1:18:00 -> 73.30h)
    - datetime época 1900 (ej: 1900-01-08 16:12:00 -> 8 días * 24 + 16.2 = 208.20h)
    - time (ej: 07:36:00 -> 7.60h)
    - string (ej: '73:18:00' -> 73.30h)
    - float serial de Excel (ej: 3.054166 * 24 = 73.30h)
    """
    if val is None:
        return 0.0

    if isinstance(val, timedelta):
        return round(val.total_seconds() / 3600.0, 2)

    if isinstance(val, datetime):
        if val.year == 1900 and val.month == 1:
            total_horas = (val.day * 24.0) + val.hour + (val.minute / 60.0) + (val.second / 3600.0)
            return round(total_horas, 2)
        return 0.0

    if isinstance(val, time):
        return round(val.hour + (val.minute / 60.0) + (val.second / 3600.0), 2)

    if isinstance(val, str):
        val_clean = val.replace("[", "").replace("]", "").strip()
        parts = val_clean.split(":")
        if len(parts) >= 2:
            try:
                h = float(parts[0])
                m = float(parts[1])
                s = float(parts[2]) if len(parts) > 2 else 0.0
                return round(h + (m / 60.0) + (s / 3600.0), 2)
            except ValueError:
                pass
        try:
            return round(float(val_clean) * 24.0, 2)
        except ValueError:
            return 0.0

    if isinstance(val, (int, float)):
        return round(float(val) * 24.0, 2)

    return 0.0


# ==============================================================================
# MÓDULO 3: GESTIÓN DE ESTADO ANTI-SPAM (Persistencia Local)
# ==============================================================================
def load_state() -> dict:
    if os.path.exists(STATE_FILE):
        try:
            with open(STATE_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def save_state(state: dict):
    try:
        with open(STATE_FILE, 'w', encoding='utf-8') as f:
            json.dump(state, f, indent=2, ensure_ascii=False)
    except Exception as e:
        logger.error(f"Error guardando estado: {e}")

def should_send_alert(plant_id: str, current_hours: float, state: dict) -> bool:
    """
    Determina si se debe enviar una alerta o silenciarla para evitar spam:
    - Si es la primera vez que supera el umbral: SÍ
    - Si ya se alertó pero las horas aumentaron >= INCREMENTO_MIN_ALERTA (ej. +6h): SÍ
    - Si pasaron más de EXPIRACION_ESTADO_HORAS (ej. 24h) como recordatorio diario: SÍ
    - De lo contrario: NO
    """
    record = state.get(plant_id)
    if not record:
        return True

    last_hours = record.get("last_hours", 0.0)
    last_timestamp = record.get("timestamp", 0)
    hours_diff = current_hours - last_hours
    time_diff_hours = (datetime.now().timestamp() - last_timestamp) / 3600.0

    if hours_diff >= INCREMENTO_MIN_ALERTA:
        return True
    if time_diff_hours >= EXPIRACION_ESTADO_HORAS:
        return True

    return False


# ==============================================================================
# MÓDULO 4: DESPACHO DE NOTIFICACIONES (Telegram y Email)
# ==============================================================================
def send_telegram_alert(chat_id: str, message: str) -> bool:
    """Envía mensaje directo por Telegram Bot (100% gratuito)."""
    if not TELEGRAM_BOT_TOKEN or not chat_id:
        return False

    url = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/sendMessage"
    payload = {
        "chat_id": chat_id,
        "text": message,
        "parse_mode": "Markdown"
    }
    try:
        res = requests.post(url, json=payload, timeout=12)
        if res.status_code == 200:
            logger.info(f"Telegram enviado a chat {chat_id}")
            return True
        else:
            logger.warning(f"Fallo Telegram HTTP {res.status_code}: {res.text}")
            return False
    except Exception as e:
        logger.error(f"Error en envío Telegram: {e}")
        return False


def send_email_alert(destinatario: str, asunto: str, html_body: str) -> bool:
    """Envía correo HTML mediante servidor SMTP gratuito."""
    if not SMTP_USER or not SMTP_PASS or not destinatario:
        return False

    msg = MIMEMultipart("alternative")
    msg["Subject"] = asunto
    msg["From"] = SMTP_FROM
    msg["To"] = destinatario
    msg.attach(MIMEText(html_body, "html"))

    try:
        with smtplib.SMTP(SMTP_SERVER, SMTP_PORT, timeout=15) as server:
            server.starttls()
            server.login(SMTP_USER, SMTP_PASS)
            server.sendmail(SMTP_FROM, destinatario, msg.as_string())
            logger.info(f"Email enviado a {destinatario}")
            return True
    except Exception as e:
        logger.error(f"Error enviando email a {destinatario}: {e}")
        return False


def build_email_template(sitio: str, estado: str, fecha_surtido: str, hora_surtido: str, horas: float, umbral: float) -> str:
    """Genera plantilla visual de alta calidad para cliente de correo."""
    exceso = round(horas - umbral, 1)
    return f"""
    <!DOCTYPE html>
    <html lang="es">
    <head>
      <meta charset="utf-8">
      <style>
        body {{ font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif; background-color: #0f172a; margin: 0; padding: 20px; color: #e2e8f0; }}
        .card {{ max-width: 580px; margin: 0 auto; background: #1e293b; border-radius: 12px; border: 1px solid #334155; overflow: hidden; box-shadow: 0 10px 25px rgba(0,0,0,0.4); }}
        .header {{ background: linear-gradient(135deg, #ef4444 0%, #b91c1c 100%); color: #ffffff; padding: 22px 24px; }}
        .badge {{ display: inline-block; background: rgba(0,0,0,0.25); border: 1px solid rgba(255,255,255,0.2); padding: 3px 8px; border-radius: 4px; font-size: 11px; font-weight: 700; text-transform: uppercase; letter-spacing: 0.05em; }}
        .header h1 {{ margin: 10px 0 2px 0; font-size: 19px; font-weight: 800; }}
        .header p {{ margin: 0; font-size: 13px; opacity: 0.9; }}
        .content {{ padding: 24px; }}
        .kpi-row {{ display: flex; gap: 12px; margin-bottom: 20px; }}
        .kpi-box {{ flex: 1; background: #0f172a; border: 1px solid #334155; border-radius: 8px; padding: 14px; text-align: center; }}
        .kpi-title {{ font-size: 11px; color: #94a3b8; font-weight: 600; text-transform: uppercase; }}
        .kpi-val {{ font-size: 26px; font-weight: 800; color: #f87171; margin-top: 4px; }}
        .info-table {{ width: 100%; border-collapse: collapse; margin-top: 10px; font-size: 13px; }}
        .info-table td {{ padding: 8px 0; border-bottom: 1px solid #334155; }}
        .info-table td.col-lbl {{ color: #94a3b8; font-weight: 500; width: 40%; }}
        .info-table td.col-val {{ color: #f8fafc; font-weight: 600; }}
        .alert-box {{ background: rgba(239, 68, 68, 0.12); border-left: 4px solid #ef4444; padding: 14px; border-radius: 4px; margin-top: 20px; font-size: 12.5px; color: #fca5a5; line-height: 1.5; }}
        .footer {{ background: #0f172a; padding: 14px 24px; font-size: 11px; color: #64748b; text-align: center; border-top: 1px solid #334155; }}
      </style>
    </head>
    <body>
      <div class="card">
        <div class="header">
          <span class="badge">Nivel de Alerta: Crítico</span>
          <h1>Consumo Excesivo en Grupo Electrógeno</h1>
          <p>Se ha superado el umbral preventivo de horas de marcha continuas.</p>
        </div>
        <div class="content">
          <div class="kpi-row">
            <div class="kpi-box">
              <div class="kpi-title">Horas Acumuladas</div>
              <div class="kpi-val">{horas:.1f} h</div>
            </div>
            <div class="kpi-box">
              <div class="kpi-title">Umbral Operativo</div>
              <div class="kpi-val" style="color: #cbd5e1;">{umbral:.0f} h <span style="font-size:12px; color:#ef4444;">(+{exceso:.1f}h)</span></div>
            </div>
          </div>
          <table class="info-table">
            <tr>
              <td class="col-lbl">Sitio / Planta:</td>
              <td class="col-val">{sitio}</td>
            </tr>
            <tr>
              <td class="col-lbl">Región / Estado:</td>
              <td class="col-val">{estado}</td>
            </tr>
            <tr>
              <td class="col-lbl">Último Surtido:</td>
              <td class="col-val">{fecha_surtido} &bull; {hora_surtido}</td>
            </tr>
          </table>
          <div class="alert-box">
            <strong>Directrices Operativas:</strong><br>
            &bull; Verificar reserva remanente de diésel para evitar descebado de bomba.<br>
            &bull; Gestionar orden prioritaria de reabastecimiento en logística.<br>
            &bull; Comprobar estado de tensión en la acometida comercial de la subestación.
          </div>
        </div>
        <div class="footer">
          Airtek Energía &bull; Sistema Automatizado de Monitoreo de Infraestructura
        </div>
      </div>
    </body>
    </html>
    """


# ==============================================================================
# MÓDULO 5: CICLO PRINCIPAL DE EJECUCIÓN
# ==============================================================================
def execute_monitoring_pipeline():
    logger.info("================================================================================")
    logger.info("INICIANDO CICLO DE MONITOREO DE GRUPOS ELECTRÓGENOS")
    logger.info("================================================================================")

    # 1. Descargar archivo en vivo desde OneDrive
    download_success = download_onedrive_excel(ONEDRIVE_SHARING_URL, LOCAL_EXCEL_FILENAME)
    if not download_success:
        if not os.path.exists(LOCAL_EXCEL_FILENAME):
            logger.error("No se pudo obtener el archivo y no existe versión local previa. Abortando.")
            return
        logger.warning("Usando copia local previa como respaldo por falla de red.")

    # 2. Cargar hoja de datos
    try:
        wb = openpyxl.load_workbook(LOCAL_EXCEL_FILENAME, data_only=True)
    except Exception as e:
        logger.error(f"Error abriendo Excel: {e}")
        return

    sheet_name = 'Lista y datos de G.E.'
    if sheet_name not in wb.sheetnames:
        logger.error(f"No se encontró la hoja '{sheet_name}'. Hojas disponibles: {wb.sheetnames}")
        return

    ws = wb[sheet_name]
    state = load_state()

    # Columnas fijas según estructura auditada del libro:
    # Col 13: FECHA DE ÚLTIMO SURTIDO
    # Col 14: HORA DE ÚLTIMO SURTIDO
    # Col 15: HORAS ENCENDIDAS DESDE ÚLTIMO SURTIDO
    # Col 16: SITIO Y NOMBRE DE GRUPO ELECTRÓGENO
    # Col 17: ESTADO
    # Si existen columnas adicionales dinámicas para CORREO o TELEGRAM, las buscamos en la fila 2
    correo_col = None
    telegram_col = None

    for c in range(1, ws.max_column + 1):
        header_val = str(ws.cell(row=2, column=c).value or '').strip().upper()
        if 'CORREO' in header_val:
            correo_col = c
        if 'TELEGRAM' in header_val or 'CHAT_ID' in header_val:
            telegram_col = c

    total_evaluadas = 0
    total_alertas = 0
    alertas_enviadas = 0

    for r in range(3, ws.max_row + 1):
        sitio_nombre = ws.cell(row=r, column=16).value
        if not sitio_nombre:
            continue

        total_evaluadas += 1
        estado = str(ws.cell(row=r, column=17).value or 'N/A').strip()
        fecha_surtido_raw = ws.cell(row=r, column=13).value
        hora_surtido_raw = ws.cell(row=r, column=14).value
        raw_horas = ws.cell(row=r, column=15).value

        # Formatear fecha y hora
        if isinstance(fecha_surtido_raw, datetime):
            fecha_surtido = fecha_surtido_raw.strftime("%d/%m/%Y")
        else:
            fecha_surtido = str(fecha_surtido_raw or "No registrado")

        if isinstance(hora_surtido_raw, time):
            hora_surtido = hora_surtido_raw.strftime("%I:%M %p")
        elif isinstance(hora_surtido_raw, datetime):
            hora_surtido = hora_surtido_raw.strftime("%I:%M %p")
        else:
            hora_surtido = str(hora_surtido_raw or "--:--")

        horas = parse_excel_duration_to_hours(raw_horas)

        # Destinatarios por celda o variables por defecto
        destinatario_correo = None
        if correo_col:
            val_correo = ws.cell(row=r, column=correo_col).value
            if val_correo:
                destinatario_correo = str(val_correo).strip()

        destinatario_telegram = TELEGRAM_DEFAULT_CHAT_ID
        if telegram_col:
            val_tg = ws.cell(row=r, column=telegram_col).value
            if val_tg:
                destinatario_telegram = str(val_tg).strip()

        # Evaluación de umbral
        if horas >= UMBRAL_HORAS_DEFAULT:
            total_alertas += 1
            plant_id = f"{sitio_nombre}_{estado}"

            if should_send_alert(plant_id, horas, state):
                logger.warning(
                    f"[ALERTA CRÍTICA] {sitio_nombre} [{estado}] - {horas:.2f} h "
                    f"(Umbral >= {UMBRAL_HORAS_DEFAULT} h | Último Surtido: {fecha_surtido} {hora_surtido})"
                )

                # Despacho Telegram Directo
                msg_tg = (
                    f"🚨 *ALERTA: GRUPO ELECTRÓGENO SOBREPASÓ LÍMITE*\n\n"
                    f"📍 *Sitio:* `{sitio_nombre}`\n"
                    f"🗺️ *Estado:* {estado}\n"
                    f"⏱️ *Horas encendidas:* `{horas:.1f} hrs` *(Límite: {UMBRAL_HORAS_DEFAULT:.0f} hrs)*\n"
                    f"⛽ *Último Surtido:* {fecha_surtido} a las {hora_surtido}\n\n"
                    f"⚠️ _Acción:_ Validar nivel de diésel y programar ruta de reabastecimiento."
                )
                tg_ok = send_telegram_alert(destinatario_telegram, msg_tg)

                # Despacho Correo Electrónico
                if destinatario_correo:
                    asunto = f"⚠️ ALERTA: {sitio_nombre} registra {horas:.1f} hrs de operación"
                    html = build_email_template(sitio_nombre, estado, fecha_surtido, hora_surtido, horas, UMBRAL_HORAS_DEFAULT)
                    send_email_alert(destinatario_correo, asunto, html)

                # Actualizar estado anti-spam
                state[plant_id] = {
                    "last_hours": horas,
                    "timestamp": datetime.now().timestamp(),
                    "date_str": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                }
                alertas_enviadas += 1
            else:
                logger.info(f"[SILENCIADO POR ANTI-SPAM] {sitio_nombre}: Ya alertado ({horas:.2f}h)")

    save_state(state)
    logger.info("--------------------------------------------------------------------------------")
    logger.info(
        f"RESUMEN: {total_evaluadas} plantas revisadas | "
        f"{total_alertas} sobre umbral (>= {UMBRAL_HORAS_DEFAULT}h) | "
        f"{alertas_enviadas} notificaciones despachadas."
    )
    logger.info("================================================================================\n")


if __name__ == '__main__':
    execute_monitoring_pipeline()
